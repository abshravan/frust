"""
Batch-process a folder of call recordings and write:
  - Excel summary (one row per file)
  - full_results.json (per-window timelines)
  - PNG timeline chart per call (score line + emotion strip)
  - dashboard.html (interactive Plotly charts; click angry points to play audio)
  - segments/<call_id>/window_<t>.wav  (angry/frustrated clips for the player)
  - InfluxDB v2 write (optional, set ENABLE_INFLUX=1)

Usage:
    python batch.py --input-dir /path/to/audio
    python batch.py --input-dir ./calls --output-dir results \\
                    --excel results.xlsx --recursive --workers 8
"""

from __future__ import annotations

import argparse
import json
import logging
import os
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from itertools import islice
from pathlib import Path

import numpy as np
import pandas as pd
from tqdm import tqdm

from audio_loader import load_audio
from dashboard import generate_dashboard
from db import write_call_to_influx
from emotion_model import MODELS_REGISTRY, get_classifier, preload_all_models
from post_processing import build_report
from segment_extractor import save_flagged_segments
from segmenter import segment_audio
from visualizer import plot_call

logger = logging.getLogger(__name__)

AUDIO_EXTENSIONS = {".wav", ".mp3", ".flac", ".ogg", ".m4a"}
_DEFAULT_WORKERS = min(8, (os.cpu_count() or 4))

EXCEL_COLUMNS = [
    "filename", "call_id", "is_frustrating", "severity",
    "frustration_score", "max_anger_score", "num_frustrating_windows",
    "num_frustration_segments", "pct_frustrated_windows",
    "longest_angry_streak", "escalation_detected",
    "total_windows", "duration_s", "processing_time_s", "reason", "status",
]


# ── Metrics helpers ───────────────────────────────────────────────────────────

def _extract_metrics(report: dict) -> dict:
    timeline = report.get("timeline", [])
    num_windows = len(timeline)
    num_angry      = sum(1 for e in timeline if e["emotion"] == "angry")
    num_frustrated = sum(1 for e in timeline if e["emotion"] == "frustrated")
    num_frustrating = num_angry + num_frustrated
    duration_s = timeline[-1]["end"] if timeline else 0.0
    longest  = _longest_angry_streak(timeline)
    num_segs = _count_angry_segments(timeline)

    return {
        "is_frustrating":           bool(report.get("flagged", False)),
        "severity":                  report.get("severity", ""),
        "frustration_score":         report.get("anger_ratio", 0.0),
        "max_anger_score":           report.get("max_anger_score", 0.0),
        "num_frustrating_windows":   num_frustrating,
        "num_frustration_segments":  num_segs,
        "pct_frustrated_windows":    round(num_frustrating / num_windows, 4)
                                     if num_windows else 0.0,
        "longest_angry_streak":      longest,
        "escalation_detected":       bool(report.get("escalation_detected", False)),
        "reason":                    report.get("reason", ""),
        "_duration_s":               round(duration_s, 2),
        "_total_windows":            num_windows,
    }


def _longest_angry_streak(timeline: list[dict]) -> int:
    longest = current = 0
    for e in timeline:
        if e["emotion"] == "angry":
            current += 1
            longest = max(longest, current)
        else:
            current = 0
    return longest


def _count_angry_segments(timeline: list[dict]) -> int:
    count, in_seg = 0, False
    for e in timeline:
        if e["emotion"] == "angry":
            if not in_seg:
                count += 1
                in_seg = True
        else:
            in_seg = False
    return count


# ── Phase 1: load + segment (I/O-bound, runs in worker threads) ──────────────

def _load_one(audio_path: Path) -> tuple[Path, np.ndarray, int, list, float]:
    """Load audio and produce segments. Returns (path, audio, sr, segments, t_start)."""
    t_start = time.time()
    audio, sr = load_audio(audio_path)
    segments = segment_audio(audio, sr)
    return audio_path, audio, sr, segments, t_start


# ── Phase 3: post-process (CPU/I/O-bound, runs in worker threads) ─────────────

def _postprocess_one(
    audio_path: Path,
    audio: np.ndarray,
    sr: int,
    predictions: list,
    t_start: float,
    chart_dirs: tuple[Path, Path],
    output_dir: Path,
) -> tuple[dict, dict]:
    filename = audio_path.name
    call_id  = audio_path.stem

    report      = build_report(predictions)
    report_dict = report.to_dict()

    dest = chart_dirs[0] if report_dict.get("flagged") else chart_dirs[1]
    try:
        plot_call(report_dict, call_id=call_id, output_path=dest)
    except Exception as exc:
        logger.warning("Chart generation failed for %s: %s", filename, exc)

    enriched = save_flagged_segments(
        audio, sr, report_dict["timeline"], call_id, output_dir
    )
    report_dict["timeline"] = enriched

    try:
        write_call_to_influx(call_id, report_dict)
    except Exception as exc:
        logger.warning("InfluxDB write skipped for %s: %s", filename, exc)

    metrics      = _extract_metrics(report_dict)
    duration_s   = metrics.pop("_duration_s", 0.0)
    total_windows = metrics.pop("_total_windows", len(predictions))

    row = {
        "filename": filename,
        "call_id":  call_id,
        **metrics,
        "total_windows":     total_windows,
        "duration_s":        duration_s,
        "processing_time_s": round(time.time() - t_start, 2),
        "status": "Success",
    }
    full = {
        "call_id":  call_id,
        "filename": filename,
        "report":   report_dict,
        "timeline": enriched,
    }
    return row, full


def _failure_row(audio_path: Path, exc: Exception) -> dict:
    row: dict = {col: "" for col in EXCEL_COLUMNS}
    row.update({
        "filename": audio_path.name,
        "call_id":  audio_path.stem,
        "status":   f"Failed: {str(exc)[:140]}",
    })
    return row


# ── Folder walker ─────────────────────────────────────────────────────────────

def _find_audio_files(root: Path, recursive: bool) -> list[Path]:
    if not root.exists():
        raise FileNotFoundError(f"Input directory not found: {root}")
    fn = root.rglob if recursive else root.glob
    found: set[Path] = set()
    for ext in AUDIO_EXTENSIONS:
        found.update(fn(f"*{ext}"))
        found.update(fn(f"*{ext.upper()}"))
    return sorted(found)


def _chunks(lst: list, n: int):
    """Yield successive n-sized chunks from lst."""
    it = iter(lst)
    while chunk := list(islice(it, n)):
        yield chunk


# ── Main batch driver ─────────────────────────────────────────────────────────

def batch_process(
    input_dir: str | Path,
    output_dir: str | Path = "results",
    excel_filename: str = "frustration_detection_results.xlsx",
    workers: int = _DEFAULT_WORKERS,
    chunk_size: int = 32,
    recursive: bool = False,
) -> pd.DataFrame:
    """
    Three-phase pipeline:

    Phase 1 – ``workers`` threads load audio and segment files in parallel.
    Phase 2 – All segments from the chunk are fed to the model in ONE batched
               GPU call (maximises GPU utilisation; no CUDA threading issues).
    Phase 3 – ``workers`` threads build reports, render charts, and save clips
               in parallel.

    ``chunk_size`` controls how many files are held in RAM at once.
    """
    input_dir  = Path(input_dir)
    output_dir = Path(output_dir)
    charts_flagged     = output_dir / "charts" / "flagged"
    charts_not_flagged = output_dir / "charts" / "not_flagged"
    for d in (output_dir, charts_flagged, charts_not_flagged):
        d.mkdir(parents=True, exist_ok=True)

    files = _find_audio_files(input_dir, recursive=recursive)
    if not files:
        raise ValueError(f"No audio files found in {input_dir}")

    logger.info(
        "Found %d audio file(s) | workers=%d | chunk_size=%d",
        len(files), workers, chunk_size,
    )
    logger.info("Pre-loading model…")
    preload_all_models()

    cfg        = MODELS_REGISTRY[0]
    classifier = get_classifier(cfg.model_id)
    chart_dirs = (charts_flagged, charts_not_flagged)

    rows: list[dict] = []
    all_full: list[dict] = []

    n_chunks = (len(files) + chunk_size - 1) // chunk_size
    for chunk_idx, chunk in enumerate(
        tqdm(_chunks(files, chunk_size), total=n_chunks, desc="Chunks", unit="chunk")
    ):
        # ── Phase 1: parallel load + segment ─────────────────────────────────
        loaded: list[tuple[Path, np.ndarray, int, list, float]] = []
        with ThreadPoolExecutor(max_workers=workers) as pool:
            futures = {pool.submit(_load_one, f): f for f in chunk}
            for fut in tqdm(
                as_completed(futures), total=len(futures),
                desc=f"  Load  chunk {chunk_idx+1}/{n_chunks}", leave=False,
            ):
                f = futures[fut]
                try:
                    loaded.append(fut.result())
                except Exception as exc:
                    logger.error("Load failed %s: %s", f.name, exc)
                    rows.append(_failure_row(f, exc))

        if not loaded:
            continue

        # ── Phase 2: single batched GPU inference ─────────────────────────────
        # Flatten all segments from all files in this chunk into one list.
        flat_segs: list = []
        slices: list[tuple[int, int]] = []
        for _path, _audio, _sr, segs, _t in loaded:
            lo = len(flat_segs)
            flat_segs.extend(segs)
            slices.append((lo, len(flat_segs)))

        logger.info(
            "Chunk %d/%d — GPU inference on %d segments from %d file(s)",
            chunk_idx + 1, n_chunks, len(flat_segs), len(loaded),
        )
        all_preds = classifier.predict_all(flat_segs)

        # ── Phase 3: parallel post-processing ────────────────────────────────
        with ThreadPoolExecutor(max_workers=workers) as pool:
            futures2: dict = {}
            for (path, audio, sr, _segs, t_start), (lo, hi) in zip(loaded, slices):
                preds = all_preds[lo:hi]
                fut = pool.submit(
                    _postprocess_one,
                    path, audio, sr, preds, t_start, chart_dirs, output_dir,
                )
                futures2[fut] = path

            for fut in tqdm(
                as_completed(futures2), total=len(futures2),
                desc=f"  Post  chunk {chunk_idx+1}/{n_chunks}", leave=False,
            ):
                f = futures2[fut]
                try:
                    row, full = fut.result()
                    rows.append(row)
                    all_full.append(full)
                except Exception as exc:
                    logger.error("Post-process failed %s: %s", f.name, exc)
                    rows.append(_failure_row(f, exc))

    df = pd.DataFrame(rows, columns=EXCEL_COLUMNS)
    df = df.sort_values("filename").reset_index(drop=True)

    excel_path    = output_dir / excel_filename
    json_path     = output_dir / "full_results.json"
    _write_excel(df, excel_path)
    with json_path.open("w") as fh:
        json.dump(all_full, fh, indent=2)

    dashboard_path = generate_dashboard(all_full, output_dir)
    _print_summary(df, excel_path, json_path, charts_flagged,
                   charts_not_flagged, dashboard_path)
    return df


# ── Excel writer ──────────────────────────────────────────────────────────────

def _write_excel(df: pd.DataFrame, path: Path) -> None:
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Summary", index=False)
        ws = writer.sheets["Summary"]

        from openpyxl.styles import Font, PatternFill
        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)

        for idx, col in enumerate(df.columns, start=1):
            cell = ws.cell(row=1, column=idx)
            cell.fill = header_fill
            cell.font = header_font
            max_len = max(len(str(col)), df[col].astype(str).str.len().max() or 0)
            ws.column_dimensions[_col_letter(idx)].width = min(max_len + 2, 52)

        alt_fill = PatternFill("solid", fgColor="D9E1F2")
        for row_idx in range(2, len(df) + 2):
            if row_idx % 2 == 0:
                for col_idx in range(1, len(df.columns) + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = alt_fill

        ws.freeze_panes = "A2"


def _col_letter(idx: int) -> str:
    letters = ""
    while idx > 0:
        idx, rem = divmod(idx - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


# ── Summary ───────────────────────────────────────────────────────────────────

def _print_summary(
    df: pd.DataFrame,
    excel_path: Path,
    json_path: Path,
    charts_flagged: Path,
    charts_not_flagged: Path,
    dashboard_path: Path,
) -> None:
    total = len(df)
    ok    = df["status"] == "Success"
    n_ok  = int(ok.sum())
    n_frustrating = int(df.loc[ok, "is_frustrating"].sum())
    avg_t = df.loc[ok, "processing_time_s"].mean() if n_ok else 0.0

    print("\n" + "=" * 64)
    print("BATCH PROCESSING COMPLETE")
    print("=" * 64)
    print(f"  Files processed  : {total}")
    print(f"  Successful       : {n_ok}")
    print(f"  Frustrating calls: {n_frustrating}/{n_ok} "
          f"({n_frustrating / max(n_ok, 1) * 100:.1f}%)")
    print(f"  Avg time / file  : {avg_t:.2f}s")
    print(f"  Excel saved      : {excel_path}")
    print(f"  JSON saved       : {json_path}")
    print(f"  Dashboard        : {dashboard_path}")
    print(f"  Charts (flagged) : {charts_flagged}/")
    print(f"  Charts (ok)      : {charts_not_flagged}/")
    print("=" * 64)


# ── CLI ───────────────────────────────────────────────────────────────────────

def _build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        description="Batch-analyze a folder of call audio and write results to Excel."
    )
    p.add_argument("--input-dir",   required=True,
                   help="Folder containing audio files")
    p.add_argument("--output-dir",  default="results",
                   help="Output folder (created if absent)")
    p.add_argument("--excel",       default="frustration_detection_results.xlsx")
    p.add_argument("--workers",     type=int, default=_DEFAULT_WORKERS,
                   help=f"Parallel threads for I/O and post-processing "
                        f"(default: {_DEFAULT_WORKERS}, auto = CPU count up to 8). "
                        f"GPU inference always runs as a single batched call.")
    p.add_argument("--chunk-size",  type=int, default=32,
                   help="Files per GPU inference chunk. Lower = less RAM, "
                        "higher = better GPU utilisation (default: 32).")
    p.add_argument("--recursive",   action="store_true",
                   help="Recurse into sub-directories")
    return p


def main() -> int:
    logging.basicConfig(
        level=os.getenv("LOG_LEVEL", "INFO").upper(),
        format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
        datefmt="%Y-%m-%dT%H:%M:%S",
    )
    args = _build_parser().parse_args()
    batch_process(
        input_dir=args.input_dir,
        output_dir=args.output_dir,
        excel_filename=args.excel,
        workers=args.workers,
        chunk_size=args.chunk_size,
        recursive=args.recursive,
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
